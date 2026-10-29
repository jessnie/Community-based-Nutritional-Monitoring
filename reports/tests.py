from datetime import date
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from nutrition.models import NutritionAssessment
from residents.models import Resident


class ReportFilteringTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_user(
			username='reportviewer',
			password='test-password',
		)
		self.client.force_login(user)

		self.ada = Resident.objects.create(
			first_name='Ada',
			last_name='Alpha',
			sex='Female',
			birth_date=date(1995, 1, 1),
			address='North Zone',
		)
		self.ben = Resident.objects.create(
			first_name='Ben',
			last_name='Beta',
			sex='Male',
			birth_date=date(1995, 1, 1),
			address='South Zone',
		)

		self.ada_match = NutritionAssessment.objects.create(
			resident=self.ada,
			assessment_date=date(2026, 4, 10),
			height=170,
			weight=75,
		)
		NutritionAssessment.objects.create(
			resident=self.ada,
			assessment_date=date(2026, 3, 10),
			height=170,
			weight=65,
		)
		NutritionAssessment.objects.create(
			resident=self.ben,
			assessment_date=date(2026, 4, 15),
			height=170,
			weight=80,
		)

	def report_filters(self):
		return {
			'status': 'Overweight',
			'resident_name': 'Ada',
			'date_from': '2026-04-01',
			'date_to': '2026-04-30',
		}

	def test_report_dashboard_applies_status_name_and_date_filters(self):
		response = self.client.get(
			reverse('report_dashboard'),
			self.report_filters(),
		)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(
			list(response.context['assessments'].values_list('pk', flat=True)),
			[self.ada_match.pk],
		)
		self.assertEqual(response.context['total_reports'], 1)
		self.assertEqual(response.context['status_summary']['Overweight'], 1)
		self.assertEqual(response.context['status_summary']['Normal'], 0)
		self.assertContains(response, 'class="status-summary-tile"')
		self.assertNotContains(response, 'Priority Alerts')

	def test_excel_export_uses_the_same_filters_as_report(self):
		response = self.client.get(reverse('export_excel'), self.report_filters())

		self.assertEqual(response.status_code, 200)
		workbook = load_workbook(BytesIO(response.content), data_only=True)
		worksheet = workbook['Nutrition Report']

		self.assertEqual(worksheet.max_row, 6)
		self.assertEqual(worksheet['B6'].value, 'Ada Alpha')

	def test_pdf_export_accepts_the_same_filters(self):
		response = self.client.get(reverse('export_pdf'), self.report_filters())

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response['Content-Type'], 'application/pdf')

	def test_invalid_resident_id_returns_no_assessments(self):
		filters = self.report_filters()
		filters['resident_id'] = 'not-an-id'

		response = self.client.get(reverse('report_dashboard'), filters)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['total_reports'], 0)

	def test_invalid_date_range_returns_no_assessments(self):
		filters = self.report_filters()
		filters['date_from'] = '2026-05-01'
		filters['date_to'] = '2026-04-01'

		response = self.client.get(reverse('report_dashboard'), filters)

		self.assertEqual(response.status_code, 200)
		self.assertTrue(response.context['date_error'])
		self.assertEqual(response.context['total_reports'], 0)
