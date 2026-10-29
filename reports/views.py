from openpyxl import Workbook
from openpyxl.styles import Font
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.shortcuts import render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.http import HttpResponse

from activity_logs.models import ActivityLog
from activity_logs.utils import log_activity

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
)
from reportlab.lib.styles import getSampleStyleSheet

from nutrition.models import NutritionAssessment, FollowUpRecord
from residents.models import Resident


def get_template(request, staff_template, admin_template):
    """
    Return the correct template depending on the logged-in user.
    """
    if request.user.is_superuser:
        return admin_template
    return staff_template


def filtered_assessments(request):
    assessments = NutritionAssessment.objects.select_related(
        'resident'
    ).order_by('-assessment_date', '-created_at')

    status = request.GET.get('status', '').strip()
    resident_name = request.GET.get('resident_name', '').strip()
    resident_id = (
        request.GET.get('resident_id', '').strip()
        or request.GET.get('resident', '').strip()
    )
    date_from_value = request.GET.get('date_from', '').strip()
    date_to_value = request.GET.get('date_to', '').strip()
    date_from = parse_date(date_from_value) if date_from_value else None
    date_to = parse_date(date_to_value) if date_to_value else None

    invalid_filter = False
    if status in dict(NutritionAssessment.STATUS_CHOICES):
        assessments = assessments.filter(nutritional_status=status)
    elif status:
        invalid_filter = True
        status = ''

    if resident_id and resident_id.isdecimal():
        assessments = assessments.filter(resident_id=resident_id)
    elif resident_id:
        invalid_filter = True

    if resident_name:
        assessments = assessments.filter(
            Q(resident__first_name__icontains=resident_name)
            | Q(resident__last_name__icontains=resident_name)
        )

    invalid_date = (
        (date_from_value and date_from is None)
        or (date_to_value and date_to is None)
    )
    invalid_range = date_from and date_to and date_from > date_to
    if invalid_date or invalid_range or invalid_filter:
        assessments = assessments.none()
    else:
        if date_from:
            assessments = assessments.filter(assessment_date__gte=date_from)
        if date_to:
            assessments = assessments.filter(assessment_date__lte=date_to)

    return assessments, {
        'status': status,
        'resident_name': resident_name,
        'resident_id': resident_id,
        'date_from': date_from_value,
        'date_to': date_to_value,
        'date_error': bool(invalid_date or invalid_range),
        'export_query': request.GET.urlencode(),
    }


@login_required(login_url='login')
def report_dashboard(request):
    assessments, filters = filtered_assessments(request)

    today = timezone.now()

    total_reports = assessments.count()

    monthly_reports = assessments.filter(
        assessment_date__month=today.month,
        assessment_date__year=today.year
    ).count()

    status_summary = {
        status_name: assessments.filter(nutritional_status=status_name).count()
        for status_name, _ in NutritionAssessment.STATUS_CHOICES
    }

    urgent_cases = assessments.filter(
        nutritional_status__in=['Underweight', 'Obese']
    ).order_by(
        '-assessment_date',
        '-bmi'
    )[:5]

    pending_followups = FollowUpRecord.objects.filter(
        assessment__in=assessments,
        status='Pending'
    ).count()

    pdf_exports = ActivityLog.objects.filter(
        action='Export PDF'
    ).count()

    excel_exports = ActivityLog.objects.filter(
        action='Export Excel'
    ).count()

    context = {
        'assessments': assessments,
        'residents': Resident.objects.order_by('last_name', 'first_name'),
        'total_reports': total_reports,
        'monthly_reports': monthly_reports,
        'status_summary': status_summary,
        'urgent_cases': urgent_cases,
        'pending_followups': pending_followups,
        'pdf_exports': pdf_exports,
        'excel_exports': excel_exports,
    }
    context.update(filters)

    return render(
        request,
        get_template(
            request,
            'reports/report_dashboard.html',
            'admin_reports/report_dashboard.html'
        ),
        context
    )


@login_required(login_url='login')
def monthly_report(request):

    today = timezone.now()

    assessments = NutritionAssessment.objects.filter(
        assessment_date__month=today.month,
        assessment_date__year=today.year
    ).select_related(
        'resident'
    ).order_by(
        '-assessment_date'
    )

    context = {
        'assessments': assessments,
        'current_month': today.strftime("%B %Y"),
    }

    return render(
        request,
        'admin_reports/monthly_report.html',
        context
    )


@login_required(login_url='login')
def annual_report(request):

    today = timezone.now()

    assessments = NutritionAssessment.objects.filter(
        assessment_date__year=today.year
    ).select_related(
        'resident'
    ).order_by(
        '-assessment_date'
    )

    context = {
        'assessments': assessments,
        'current_year': today.year,
    }

    return render(
        request,
        'admin_reports/annual_report.html',
        context
    )


@login_required(login_url='login')
def export_pdf(request):

    response = HttpResponse(
        content_type='application/pdf'
    )

    response['Content-Disposition'] = (
        'attachment; filename="Nutrition_Report.pdf"'
    )

    doc = SimpleDocTemplate(
        response,
        pagesize=letter
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b>Community-Based Nutritional Monitoring and Decision Support System</b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            "Nutrition Assessment Report",
            styles["Heading2"]
        )
    )

    elements.append(
        Paragraph(
            f"Generated on: {timezone.now().strftime('%B %d, %Y %I:%M %p')}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph("<br/>", styles["Normal"])
    )

    assessments, _ = filtered_assessments(request)

    data = [[
        "Date",
        "Resident",
        "Age",
        "Sex",
        "Height",
        "Weight",
        "BMI",
        "Status",
        "Remarks"
    ]]

    for assessment in assessments:

        data.append([

            str(assessment.assessment_date),

            f"{assessment.resident.first_name} {assessment.resident.last_name}",

            str(assessment.resident.age),

            assessment.resident.sex,

            f"{assessment.height} cm",

            f"{assessment.weight} kg",

            str(assessment.bmi),

            assessment.nutritional_status,

            assessment.remarks or ""

        ])

    table = Table(data)

    table.setStyle(

        TableStyle([

            ('BACKGROUND', (0, 0), (-1, 0), colors.green),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),

            ('GRID', (0, 0), (-1, -1), 1, colors.black),

            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),

            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),

            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),

            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),

        ])

    )

    elements.append(table)

    doc.build(elements)

    log_activity(
        request.user,
        'Export PDF',
        'Exported nutrition report as PDF'
    )

    return response


@login_required(login_url='login')
def export_excel(request):

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename="Nutrition_Report.xlsx"'

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Nutrition Report"

    worksheet['A1'] = "Community-Based Nutritional Monitoring and Decision Support System"
    worksheet['A2'] = "Nutrition Assessment Report"
    worksheet['A3'] = f"Generated: {timezone.now().strftime('%B %d, %Y %I:%M %p')}"

    headers = [
        "Date",
        "Resident",
        "Age",
        "Sex",
        "Height (cm)",
        "Weight (kg)",
        "BMI",
        "Nutritional Status",
        "Remarks",
    ]

    row = 5

    for col_num, header in enumerate(headers, 1):

        cell = worksheet.cell(
            row=row,
            column=col_num
        )

        cell.value = header
        cell.font = Font(bold=True)

    assessments, _ = filtered_assessments(request)

    row = 6

    for assessment in assessments:

        worksheet.cell(
            row=row,
            column=1
        ).value = str(assessment.assessment_date)

        worksheet.cell(
            row=row,
            column=2
        ).value = f"{assessment.resident.first_name} {assessment.resident.last_name}"

        worksheet.cell(
            row=row,
            column=3
        ).value = assessment.resident.age

        worksheet.cell(
            row=row,
            column=4
        ).value = assessment.resident.sex

        worksheet.cell(
            row=row,
            column=5
        ).value = float(assessment.height)

        worksheet.cell(
            row=row,
            column=6
        ).value = float(assessment.weight)

        worksheet.cell(
            row=row,
            column=7
        ).value = float(assessment.bmi)

        worksheet.cell(
            row=row,
            column=8
        ).value = assessment.nutritional_status

        worksheet.cell(
            row=row,
            column=9
        ).value = assessment.remarks

        row += 1

    # Adjust column widths
    column_widths = {
        'A': 15,
        'B': 28,
        'C': 8,
        'D': 12,
        'E': 15,
        'F': 15,
        'G': 10,
        'H': 20,
        'I': 45,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    workbook.save(response)

    log_activity(
        request.user,
        'Export Excel',
        'Exported nutrition report as Excel'
    )

    return response